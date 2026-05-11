# -*- coding: utf-8 -*-
"""
资信标自动填充系统 - 完整功能验证测试
按照技术路线文档验证所有核心功能
"""

import os
import sys
import re
import json
import time
from pathlib import Path
from datetime import datetime
from collections import deque

# 设置控制台编码
if sys.platform == 'win32':
    import io
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8')

# 添加项目路径
sys.path.insert(0, str(Path(__file__).parent))

from docx import Document
from docx.document import Document as _Document
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from dotenv import load_dotenv
from openai import OpenAI

# 加载环境变量
load_dotenv()

# 配置
KIMI_API_KEY = os.getenv("KIMI_API_KEY", "")
KIMI_BASE_URL = os.getenv("KIMI_BASE_URL", "https://api.moonshot.cn/v1")
KIMI_MODEL = os.getenv("KIMI_MODEL", "moonshot-v1-32k")

# 测试文件路径
TEMPLATE_PATH = r"D:\桌面\投标文件组成模版01.docx"
COMPANY_INFO_PATH = r"D:\桌面\企业信息01.txt"
OUTPUT_PATH = r"D:\桌面\测试输出_完整填充结果_v3.docx"

print("=" * 70)
print("资信标自动填充系统 - 完整功能验证测试")
print("=" * 70)
print(f"开始时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

# 初始化客户端
client = OpenAI(api_key=KIMI_API_KEY, base_url=KIMI_BASE_URL)


# ============== 核心功能函数 ==============

def get_run_format(run):
    font = run.font
    return (
        font.name,
        font.size.pt if font.size else None,
        font.bold,
        font.italic,
        font.underline,
        run._element.rPr
    )


def merge_adjacent_runs(doc: Document) -> int:
    """合并格式相同的相邻Run"""
    merged_count = 0

    for para in doc.paragraphs:
        merged_count += _merge_runs_in_para(para)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    merged_count += _merge_runs_in_para(para)

    return merged_count


def _merge_runs_in_para(para) -> int:
    merged_count = 0
    if len(para.runs) <= 1:
        return merged_count

    i = 0
    while i < len(para.runs) - 1:
        current_run = para.runs[i]
        next_run = para.runs[i + 1]
        if get_run_format(current_run) == get_run_format(next_run):
            current_run.text += next_run.text
            next_run._element.getparent().remove(next_run._element)
            merged_count += 1
        else:
            i += 1

    return merged_count


def replace_text_in_paragraph(para, old_text: str, new_text) -> bool:
    """跨Run替换文本，保留格式"""
    if not old_text:
        return False

    new_text = str(new_text) if new_text is not None else ""
    full_text = para.text

    if old_text not in full_text:
        return False

    start_idx = full_text.find(old_text)
    end_idx = start_idx + len(old_text)

    run_start_idx = 0
    matching_runs = []

    for run in para.runs:
        run_end_idx = run_start_idx + len(run.text)
        if (run_start_idx <= start_idx < run_end_idx or
            run_start_idx < end_idx <= run_end_idx or
            (start_idx <= run_start_idx and run_end_idx <= end_idx)):
            matching_runs.append({
                'run': run,
                'run_start': run_start_idx,
                'run_end': run_end_idx,
                'text_start': max(start_idx, run_start_idx),
                'text_end': min(end_idx, run_end_idx)
            })
        run_start_idx = run_end_idx

    if not matching_runs:
        return False

    if len(matching_runs) == 1:
        run_info = matching_runs[0]
        run = run_info['run']
        run_local_start = run_info['text_start'] - run_info['run_start']
        run_local_end = run_info['text_end'] - run_info['run_start']
        run.text = run.text[:run_local_start] + new_text + run.text[run_local_end:]
        return True

    first_run_info = matching_runs[0]
    first_run = first_run_info['run']
    first_run_local_start = first_run_info['text_start'] - first_run_info['run_start']

    new_run_text = first_run.text[:first_run_local_start] + new_text
    first_run.text = new_run_text

    for i, run_info in enumerate(matching_runs):
        if i == 0:
            continue
        run = run_info['run']
        run_local_end = run_info['text_end'] - run_info['run_start']

        if run_local_end < len(run.text):
            run.text = run.text[run_local_end:]
        else:
            run.text = ""

    return True


def iter_block_items(parent):
    """按文档物理顺序遍历"""
    if isinstance(parent, _Document):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        raise ValueError("不支持的文档对象")

    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def get_tables_with_upper_context(doc: Document, window_size: int = 5) -> list:
    """获取带上下文的表格列表"""
    upper_context_window = deque(maxlen=window_size)
    table_context_list = []

    for block in iter_block_items(doc):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if text:
                upper_context_window.append(text)

        elif isinstance(block, Table):
            table_upper_context = "\n".join(upper_context_window)
            serialized_table, unique_cells_info = serialize_table(block)

            table_context_list.append({
                "table": block,
                "upper_context": table_upper_context,
                "serialized_table": serialized_table,
                "unique_cells_info": unique_cells_info
            })

            upper_context_window.clear()

    return table_context_list


def is_blank_cell(text: str) -> bool:
    """判断单元格是否为空白"""
    if not text or not text.strip():
        return True

    blank_patterns = [
        r'^_{2,}$',
        r'^\s*$',
        r'^\[空\]$',
        r'^待填写$',
        r'^待补全$',
        r'^□+$',
    ]

    text_clean = text.strip()
    for pattern in blank_patterns:
        if re.match(pattern, text_clean):
            return True

    cleaned = re.sub(r'[\s\_\-:\(\)（）\[\]【】□]', '', text_clean)
    if not cleaned:
        return True

    return False


def serialize_table(table) -> tuple:
    """序列化表格为 | 分隔符格式"""
    rows_text = []
    unique_cells_info = []
    seen_cell_ids = set()

    for row_idx, row in enumerate(table.rows):
        cells_text = []
        row_cells_info = []

        for cell_idx, cell in enumerate(row.cells):
            cell_id = id(cell._tc)

            if cell_id not in seen_cell_ids:
                seen_cell_ids.add(cell_id)
                text = cell.text.strip()

                if is_blank_cell(text):
                    cells_text.append('[空]')
                    row_cells_info.append({
                        'row_idx': row_idx,
                        'cell_idx': cell_idx,
                        'cell': cell,
                        'original_text': '',
                        'is_empty': True
                    })
                else:
                    text = text.replace('|', '｜').replace('\n', ' ')
                    cells_text.append(text)
                    row_cells_info.append({
                        'row_idx': row_idx,
                        'cell_idx': cell_idx,
                        'cell': cell,
                        'original_text': text,
                        'is_empty': False
                    })
            else:
                cells_text.append('[合并]')
                row_cells_info.append({
                    'row_idx': row_idx,
                    'cell_idx': cell_idx,
                    'cell': cell,
                    'is_merged': True
                })

        rows_text.append('| ' + ' | '.join(cells_text) + ' |')
        unique_cells_info.append(row_cells_info)

    return '\n'.join(rows_text), unique_cells_info


def detect_dynamic_list(table_context: dict) -> tuple:
    """检测是否为动态列表"""
    serialized = table_context['serialized_table']
    rows = parse_serialized_table(serialized)

    if len(rows) < 3:
        return False, None

    empty_streak = 0
    max_streak = 0

    for i, row in enumerate(rows):
        empty_count = sum(1 for cell in row if cell in ['[空]', '[合并]'])
        total_cells = len(row)

        if total_cells > 0 and empty_count / total_cells > 0.7:
            empty_streak += 1
            max_streak = max(max_streak, empty_streak)
        else:
            empty_streak = 0

    return max_streak >= 3, None


def parse_serialized_table(text: str) -> list:
    """解析 | 分隔符表格"""
    rows = []
    for line in text.strip().split('\n'):
        line = line.strip()
        if not line:
            continue
        if '|' in line:
            cells = [c.strip() for c in line.split('|')]
            cells = [c for c in cells if c]
            if cells:
                rows.append(cells)
    return rows


def call_kimi_api(prompt: str, system_prompt: str, max_tokens: int = 2048) -> str:
    """通用API调用"""
    response = client.chat.completions.create(
        model=KIMI_MODEL,
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": prompt}
        ],
        temperature=0.1,
        max_tokens=max_tokens
    )
    return response.choices[0].message.content.strip()


def call_kimi_tag_paragraphs(text_block: str) -> str:
    """为段落打标"""
    prompt = f"""【角色】你是一名专业的投标助理，精通各类投标文件模板的语义理解。
【任务】识别文本中所有需要投标人填写、补充或提供的空白处。

【空白处类型包括但不限于】:
1. 下划线:________
2. 括号提示:(项目名称)、[请填写]、（单位盖章）
3. 散落空格:  年  月  日
4. 冒号留白:投标人:

【打标规则】:
1. 根据上下文语义，为每个空白处生成一个简洁的中文标签
2. 标签格式统一为 {{{{标签名}}}}
3. 必须返回原句，只替换空白处为标签
4. 禁止生成泛标签如{{{{姓名}}}}、{{{{地址}}}}
5. 完整保留[数字]索引

【待打标文本块】
{text_block}"""

    system_prompt = "你是标签生成器，只输出打标后的内容。"
    return call_kimi_api(prompt, system_prompt, max_tokens=4096)


def call_kimi_tag_table(serialized_table: str, upper_context: str, is_dynamic: bool) -> str:
    """为表格打标"""
    if is_dynamic:
        prompt = f"""【角色】标书模板表格智能打标专家

【任务】这是动态列表表格。仅保留表头和第一行空白位，删除后续所有重复空行。

【输入】
=== 表格前文背景 ===
{upper_context}
=== 待打标表格 ===
{serialized_table}

【打标规则】
1. 在首行添加：{{{{tr for item in 列表名称}}}}
2. 使用泛型标签：{{{{item.字段名}}}}
3. 列表名称根据上下文命名

【输出】
只输出打标后的表格文本"""
    else:
        prompt = f"""【角色】标书模板表格智能打标专家

【任务】这是静态表单。所有[空]逐一替换为独立业务标签。

【输入】
=== 表格前文背景 ===
{upper_context}
=== 待打标表格 ===
{serialized_table}

【打标规则】
1. 标签名必须高度具体
2. 保留所有行列结构

【输出】
只输出打标后的表格文本"""

    system_prompt = "你是表格打标工具，只输出打标后的表格。"
    return call_kimi_api(prompt, system_prompt, max_tokens=4096)


def call_kimi_match_info(tags: list, company_info: str) -> dict:
    """信息匹配"""
    prompt = f"""【角色】企业信息匹配专家

【标签列表】：
{json.dumps(tags, ensure_ascii=False, indent=2)}

【企业信息】：
{company_info[:4000]}

【输出要求】
返回JSON对象，找不到的填"待补全"。支持列表类型。"""

    system_prompt = "你是数据匹配工具，只输出JSON。"

    result = call_kimi_api(prompt, system_prompt, max_tokens=8192)

    # 解析JSON
    result = re.sub(r'^```json\s*', '', result)
    result = re.sub(r'^```\s*', '', result)
    result = re.sub(r'```\s*$', '', result)

    json_match = re.search(r'\{[\s\S]*\}', result)
    if json_match:
        result = json_match.group()

    return json.loads(result)


def extract_tags_from_doc(doc: Document) -> list:
    """提取所有标签"""
    tags = set()
    pattern = r'\{\{(.*?)\}\}'

    for para in doc.paragraphs:
        matches = re.findall(pattern, para.text)
        tags.update(matches)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    matches = re.findall(pattern, para.text)
                    tags.update(matches)

    return sorted(list(tags))


def parse_tagged_paragraphs(text: str) -> dict:
    """解析带索引的打标结果"""
    line_map = {}
    pattern = r'^\[(\d+)\]\s*(.*)$'

    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        m = re.match(pattern, line)
        if m:
            idx = int(m.group(1))
            content = m.group(2)
            line_map[idx] = content

    return line_map


# ============== 主测试流程 ==============

print("\n" + "=" * 70)
print("【Step 1】加载文档和配置")
print("-" * 70)

doc = Document(TEMPLATE_PATH)
print(f"[OK] 文档加载成功")
print(f"  - 段落数量: {len(doc.paragraphs)}")
print(f"  - 表格数量: {len(doc.tables)}")

with open(COMPANY_INFO_PATH, 'r', encoding='utf-8') as f:
    company_info = f.read()
print(f"[OK] 企业信息已加载，共 {len(company_info)} 字符")

print("\n" + "=" * 70)
print("【Step 2】合并碎片化Run")
print("-" * 70)

start_time = time.time()
merged = merge_adjacent_runs(doc)
elapsed = time.time() - start_time
print(f"[OK] Run合并完成，共合并 {merged} 个，耗时 {elapsed:.2f}秒")

print("\n" + "=" * 70)
print("【Step 3】表格上下文增强测试")
print("-" * 70)

tables_with_context = get_tables_with_upper_context(doc)
print(f"[OK] 获取到 {len(tables_with_context)} 个带上下文的表格")

for i, table_ctx in enumerate(tables_with_context[:3]):
    print(f"\n  表格{i+1}:")
    print(f"    上文: {table_ctx['upper_context'][:50]}...")
    is_dynamic, _ = detect_dynamic_list(table_ctx)
    print(f"    动态列表: {'是' if is_dynamic else '否'}")
    print(f"    序列化预览: {table_ctx['serialized_table'][:100]}...")

print("\n" + "=" * 70)
print("【Step 4】AI辅助打标 - 段落测试")
print("-" * 70)

# 测试段落打标（只处理前5个段落）
valid_paragraphs = [p for p in doc.paragraphs if p.text.strip()]
test_paras = valid_paragraphs[:5]

for i, para in enumerate(test_paras):
    text = para.text.strip()
    if text:
        print(f"\n  测试段落{i+1}:")
        print(f"    原文: {text[:60]}...")

        try:
            test_block = f"[0] {text}"
            tagged = call_kimi_tag_paragraphs(test_block)
            print(f"    打标: {tagged[:80]}...")
        except Exception as e:
            print(f"    [WARN] 打标失败: {e}")

print("\n" + "=" * 70)
print("【Step 5】AI辅助打标 - 表格测试")
print("-" * 70)

# 测试第一个表格
if tables_with_context:
    table_ctx = tables_with_context[0]
    is_dynamic = detect_dynamic_list(table_ctx)[0]

    print(f"\n  测试表格1:")
    print(f"    动态列表: {'是' if is_dynamic else '否'}")

    try:
        tagged = call_kimi_tag_table(
            table_ctx['serialized_table'][:300],
            table_ctx['upper_context'][:100],
            is_dynamic
        )
        print(f"    打标预览: {tagged[:150]}...")
    except Exception as e:
        print(f"    [WARN] 打标失败: {e}")

print("\n" + "=" * 70)
print("【Step 6】完整流程测试（限制处理数量）")
print("-" * 70)

# 重新加载文档
doc = Document(TEMPLATE_PATH)
merge_adjacent_runs(doc)

# 只处理前10个段落和前3个表格（节省API调用）
print("\n处理段落（前10个）...")
valid_paragraphs = [p for p in doc.paragraphs if p.text.strip()][:10]
paragraph_processed = 0

for i, para in enumerate(valid_paragraphs):
    text = para.text.strip()
    if text and len(text) < 500:
        try:
            test_block = f"[{i}] {text}"
            tagged = call_kimi_tag_paragraphs(test_block)
            tagged_map = parse_tagged_paragraphs(tagged)

            if i in tagged_map:
                new_text = tagged_map[i]
                if "{{" in new_text and new_text != text:
                    replace_text_in_paragraph(para, text, new_text)
                    paragraph_processed += 1
                    print(f"  段落{i+1}: 打标成功")
        except Exception as e:
            print(f"  段落{i+1}: 失败 - {e}")

print(f"\n处理表格（前3个）...")
tables_with_context = get_tables_with_upper_context(doc)[:3]
table_processed = 0

for table_idx, table_ctx in enumerate(tables_with_context):
    table = table_ctx['table']
    is_dynamic = detect_dynamic_list(table_ctx)[0]

    try:
        tagged = call_kimi_tag_table(
            table_ctx['serialized_table'],
            table_ctx['upper_context'],
            is_dynamic
        )

        tagged_rows = parse_serialized_table(tagged)

        # 简单应用（不处理动态列表删除行）
        for row_idx, row in enumerate(table.rows[:len(tagged_rows)]):
            if row_idx >= len(tagged_rows):
                break
            tagged_cells = tagged_rows[row_idx]
            for cell_idx, cell in enumerate(row.cells[:len(tagged_cells)]):
                if cell_idx >= len(tagged_cells):
                    break
                tagged_text = tagged_cells[cell_idx]
                if '{{' in tagged_text:
                    for para in cell.paragraphs:
                        if para.text.strip():
                            replace_text_in_paragraph(para, para.text, tagged_text)

        table_processed += 1
        print(f"  表格{table_idx+1}: 处理完成")
    except Exception as e:
        print(f"  表格{table_idx+1}: 失败 - {e}")

print("\n" + "=" * 70)
print("【Step 7】标签提取")
print("-" * 70)

tags = extract_tags_from_doc(doc)
print(f"[OK] 提取到 {len(tags)} 个标签")
print(f"  标签列表: {tags[:20]}{'...' if len(tags) > 20 else ''}")

print("\n" + "=" * 70)
print("【Step 8】企业信息匹配")
print("-" * 70)

start_time = time.time()
try:
    matched_data = call_kimi_match_info(tags[:30], company_info)  # 只匹配前30个
    elapsed = time.time() - start_time
    print(f"[OK] 信息匹配完成，耗时 {elapsed:.2f}秒")
    print(f"  匹配结果数量: {len(matched_data)}")

    # 显示部分结果
    print("\n  匹配结果预览:")
    for i, (k, v) in enumerate(list(matched_data.items())[:10]):
        v_preview = str(v)[:30] if isinstance(v, str) else str(v)[:50]
        print(f"    {k}: {v_preview}")
except Exception as e:
    print(f"[WARN] 信息匹配失败: {e}")
    matched_data = {}

print("\n" + "=" * 70)
print("【Step 9】内容填充")
print("-" * 70)

replacement_count = 0

# 处理正文
for para in doc.paragraphs:
    for tag, value in matched_data.items():
        placeholder = f"{{{{{tag}}}}}"
        if placeholder in para.text:
            if replace_text_in_paragraph(para, placeholder, value):
                replacement_count += 1

# 处理表格
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for tag, value in matched_data.items():
                    placeholder = f"{{{{{tag}}}}}"
                    if placeholder in para.text:
                        if replace_text_in_paragraph(para, placeholder, value):
                            replacement_count += 1

print(f"[OK] 内容填充完成，共替换 {replacement_count} 处")

print("\n" + "=" * 70)
print("【Step 10】保存结果")
print("-" * 70)

doc.save(OUTPUT_PATH)
print(f"[OK] 文档已保存至: {OUTPUT_PATH}")

print("\n" + "=" * 70)
print("完整功能验证测试完成!")
print("=" * 70)
print(f"结束时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
print(f"\n测试结果摘要:")
print(f"  - 段落打标: {paragraph_processed} 个")
print(f"  - 表格打标: {table_processed} 个")
print(f"  - 标签提取: {len(tags)} 个")
print(f"  - 信息匹配: {len(matched_data)} 个")
print(f"  - 内容填充: {replacement_count} 处替换")
print(f"\n输出文件: {OUTPUT_PATH}")
print("\n请打开输出文件检查格式保留效果！")
