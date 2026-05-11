# -*- coding: utf-8 -*-
"""
检查生成的标签格式
"""

import os
import sys
import re
from pathlib import Path

if sys.platform == 'win32':
    import io
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from docx import Document

OUTPUT_PATH = r"D:\桌面\最终输出_完整填充_v6.docx"

doc = Document(OUTPUT_PATH)

print("=" * 70)
print("检查生成文档中的标签格式")
print("=" * 70)

# 检查段落
print("\n【段落中的标签】")
for i, para in enumerate(doc.paragraphs):
    text = para.text
    if '{%' in text or '{{' in text:
        print(f"\n段落 {i}: {text[:100]}...")
        # 找出所有标签
        jinja_tags = re.findall(r'\{[%{].*?\[%}]\}', text)
        double_tags = re.findall(r'\{\{.*?\}\}', text)
        if jinja_tags:
            print(f"  Jinja标签: {jinja_tags}")
        if double_tags:
            print(f"  双花括号: {double_tags[:5]}")

# 检查表格
print("\n【表格中的标签】")
for table_idx, table in enumerate(doc.tables):
    print(f"\n表格 {table_idx + 1}:")
    for row_idx, row in enumerate(table.rows[:5]):
        for cell_idx, cell in enumerate(row.cells):
            text = cell.text.strip()
            if '{%' in text or '{{' in text:
                print(f"  行{row_idx}列{cell_idx}: {text[:80]}...")
                # 检查标签格式
                if '{%' in text:
                    print(f"    [注意] 包含Jinja控制标签")
                if '{{' in text:
                    tags = re.findall(r'\{\{.*?\}\}', text)
                    print(f"    双花括号标签: {tags[:3]}")