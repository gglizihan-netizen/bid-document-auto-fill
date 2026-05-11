# -*- coding: utf-8 -*-
"""
资信标自动填充系统 - 完整解决方案（使用docxtpl处理动态表格）
"""

import os
import sys
import re
import json
import time
from pathlib import Path
from datetime import datetime
from collections import deque

if sys.platform == 'win32':
    import io
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8')

sys.path.insert(0, str(Path(__file__).parent))

from docx import Document
from docx.document import Document as _Document
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docxtpl import DocxTemplate
from dotenv import load_dotenv
from openai import OpenAI

load_dotenv()

KIMI_API_KEY = os.getenv("KIMI_API_KEY", "")
KIMI_BASE_URL = os.getenv("KIMI_BASE_URL", "https://api.moonshot.cn/v1")
KIMI_MODEL = os.getenv("KIMI_MODEL", "moonshot-v1-32k")

TEMPLATE_PATH = r"D:\桌面\投标文件组成模版01.docx"
COMPANY_INFO_PATH = r"D:\桌面\企业信息01.txt"
OUTPUT_PATH = r"D:\桌面\最终输出_完整填充_v5.docx"

print("=" * 70)
print("资信标自动填充系统 - 完整解决方案（支持动态表格）")
print("=" * 70)
print(f"开始时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

client = OpenAI(api_key=KIMI_API_KEY, base_url=KIMI_BASE_URL)


# ============== 辅助函数 ==============

def get_run_format(run):
    font = run.font
    return (
        font.name,
        font.size.pt if font.size else None,
        font.bold,
        font.italic,
        font.underline,
        run._element.rPr
    )


def merge_adjacent_runs(doc: Document) -> int:
    merged_count = 0
    for para in doc.paragraphs:
        merged_count += _merge_runs_in_para(para)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    merged_count += _merge_runs_in_para(para)
    return merged_count


def _merge_runs_in_para(para) -> int:
    merged_count = 0
    if len(para.runs) <= 1:
        return merged_count
    i = 0
    while i < len(para.runs) - 1:
        current_run = para.runs[i]
        next_run = para.runs[i + 1]
        if get_run_format(current_run) == get_run_format(next_run):
            current_run.text += next_run.text
            next_run._element.getparent().remove(next_run._element)
            merged_count += 1
        else:
            i += 1
    return merged_count


def iter_block_items(parent):
    if isinstance(parent, _Document):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        raise ValueError("不支持的文档对象")
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def get_tables_with_upper_context(doc: Document, window_size: int = 5) -> list:
    upper_context_window = deque(maxlen=window_size)
    table_context_list = []
    for block in iter_block_items(doc):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if text:
                upper_context_window.append(text)
        elif isinstance(block, Table):
            table_upper_context = "\n".join(upper_context_window)
            serialized_table, unique_cells_info = serialize_table(block)
            table_context_list.append({
                "table": block,
                "upper_context": table_upper_context,
                "serialized_table": serialized_table,
                "unique_cells_info": unique_cells_info
            })
            upper_context_window.clear()
    return table_context_list


def is_blank_cell(text: str) -> bool:
    if not text or not text.strip():
        return True
    text = text.strip()
    if text in ['[空]', '[合并]', '待填写', '待补全']:
        return True
    if re.match(r'^_{2,}$', text):
        return True
    cleaned = re.sub(r'[\s\_\-:\(\)（）\[\]【】]', '', text)
    if not cleaned:
        return True
    return False


def serialize_table(table) -> tuple:
    rows_text = []
    unique_cells_info = []
    seen_cell_ids = set()
    for row_idx, row in enumerate(table.rows):
        cells_text = []
        row_cells_info = []
        for cell_idx, cell in enumerate(row.cells):
            cell_id = id(cell._tc)
            if cell_id not in seen_cell_ids:
                seen_cell_ids.add(cell_id)
                text = cell.text.strip()
                if is_blank_cell(text):
                    cells_text.append('[空]')
                    row_cells_info.append({
                        'row_idx': row_idx,
                        'cell_idx': cell_idx,
                        'cell': cell,
                        'original_text': '',
                        'is_empty': True
                    })
                else:
                    text = text.replace('|', '｜').replace('\n', ' ')
                    cells_text.append(text)
                    row_cells_info.append({
                        'row_idx': row_idx,
                        'cell_idx': cell_idx,
                        'cell': cell,
                        'original_text': text,
                        'is_empty': False
                    })
            else:
                cells_text.append('[合并]')
                row_cells_info.append({
                    'row_idx': row_idx,
                    'cell_idx': cell_idx,
                    'cell': cell,
                    'is_merged': True
                })
        rows_text.append('| ' + ' | '.join(cells_text) + ' |')
        unique_cells_info.append(row_cells_info)
    return '\n'.join(rows_text), unique_cells_info


def parse_serialized_table(text: str) -> list:
    rows = []
    for line in text.strip().split('\n'):
        line = line.strip()
        if not line:
            continue
        if '|' in line:
            cells = [c.strip() for c in line.split('|')]
            cells = [c for c in cells if c]
            if cells:
                rows.append(cells)
    return rows


def detect_dynamic_list(table_context: dict) -> bool:
    serialized = table_context['serialized_table']
    rows = parse_serialized_table(serialized)
    if len(rows) < 3:
        return False
    empty_streak = 0
    max_streak = 0
    for row in rows:
        empty_count = sum(1 for cell in row if cell in ['[空]', '[合并]'])
        total_cells = len(row)
        if total_cells > 0 and empty_count / total_cells > 0.7:
            empty_streak += 1
            max_streak = max(max_streak, empty_streak)
        else:
            empty_streak = 0
    return max_streak >= 3


def call_kimi_api(prompt: str, system_prompt: str, max_tokens: int = 4096) -> str:
    response = client.chat.completions.create(
        model=KIMI_MODEL,
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": prompt}
        ],
        temperature=0.1,
        max_tokens=max_tokens
    )
    return response.choices[0].message.content.strip()


def call_kimi_tag_paragraphs_batch(paragraphs: list) -> dict:
    """批量处理段落打标"""
    indexed_text = "\n".join([f"[{i}] {p.text.strip()}" for i, p in paragraphs if p.text.strip()])

    prompt = f"""【角色】你是专业的投标文件空白处识别专家。

【任务】识别以下文本中所有需要填写的内容，并生成对应标签。

【空白处类型】：
1. 下划线：________、____、_等
2. 括号提示：（项目名称）、（单位盖章）、(招标人名称)、[请填写]等
3. 多个空格：连续3个以上空格表示的空白处
4. 冒号后留白：投标人：   （冒号后无内容或只有空格）
5. 日期格式：  年  月  日（需要填写年月日的位置）

【打标规则】：
1. 标签格式：{{{{标签名}}}}
2. 标签名必须高度具体，结合上下文
3. 必须完整保留原文所有内容，只替换空白处
4. 必须完整保留每行开头的[数字]索引
5. 如果某行没有空白处，原样返回该行

【标签命名示例】：
- "（招标人名称）" -> "{{{{招标人名称}}}}"
- "投标人：" 后空白 -> "{{{{投标单位名称}}}}"
- "  年  月  日" -> "{{{{年}}}}年{{{{月}}}}月{{{{日}}}}日"
- "（单位盖章）" -> "{{{{单位盖章}}}}"
- "联系人：" 后空白 -> "{{{{联系人姓名}}}}"
- "电话：" 后空白 -> "{{{{联系电话}}}}"
- "地址：" 后空白 -> "{{{{联系地址}}}}"

【待处理文本】：
{indexed_text[:8000]}

【输出要求】：
只输出打标后的文本，完整保留所有[数字]索引行！"""

    system_prompt = "你是标签生成器，只输出打标后的内容。"

    try:
        result = call_kimi_api(prompt, system_prompt, max_tokens=8192)
        tagged_map = {}
        pattern = r'^\[(\d+)\]\s*(.*)$'
        for line in result.splitlines():
            line = line.strip()
            if not line:
                continue
            m = re.match(pattern, line)
            if m:
                idx = int(m.group(1))
                content = m.group(2)
                tagged_map[idx] = content
        return tagged_map
    except Exception as e:
        print(f"  [ERROR] 批量打标失败: {e}")
        return {}


def call_kimi_tag_table(serialized_table: str, upper_context: str, is_dynamic: bool) -> str:
    """调用Kimi API为表格打标"""
    if is_dynamic:
        prompt = f"""【角色】动态列表表格打标专家

【表格前文背景】：
{upper_context}

【待打标表格】：
{serialized_table}

【任务】这是动态列表表格（人员表/业绩表等），请：
1. 仅保留表头行和第一个数据行
2. 删除其他所有空白数据行
3. 在第一个数据行开头添加：{{{{tr for item in 列表名称}}}}
4. 单元格使用泛型标签：{{{{item.字段名}}}}
5. 列表名称根据上下文命名（如"项目负责人列表"、"项目管理班子列表"、"业绩列表"）

【重要】第一数据行必须以{{{{tr开头，如：{{{{tr for item in 项目负责人列表}}{{{{item.姓名}}}}"""

    else:
        prompt = f"""【角色】静态表单打标专家

【表格前文背景】：
{upper_context}

【待打标表格】：
{serialized_table}

【任务】这是静态表单，请：
1. 将所有[空]替换为具体标签{{{{字段名}}}}
2. 保留所有行列结构
3. 标签名必须结合上下文，高度具体
4. 禁止使用泛标签"""

    system_prompt = "你是表格打标工具，只输出打标后的表格。"
    return call_kimi_api(prompt, system_prompt, max_tokens=8192)


def call_kimi_match_info(tags: list, company_info: str) -> dict:
    """信息匹配（支持列表类型）"""
    if not tags:
        return {}

    # 分类标签
    simple_tags = []
    list_tags = []

    for tag in tags:
        if tag.startswith('item.'):
            continue  # item.xxx 标签在循环中处理
        elif '列表' in tag or tag.startswith('tr for'):
            list_tags.append(tag)
        else:
            simple_tags.append(tag)

    all_results = {}

    # 处理简单标签
    if simple_tags:
        batch_size = 80
        for batch_start in range(0, len(simple_tags), batch_size):
            batch_tags = simple_tags[batch_start:batch_start + batch_size]

            prompt = f"""【角色】企业信息匹配专家

【标签列表】：
{json.dumps(batch_tags, ensure_ascii=False, indent=2)}

【企业信息】：
{company_info[:5000]}

【任务】为每个标签找到对应值。找不到填"待补全"。

【字段映射】：
- 投标单位名称/投标人名称/单位名称 -> 企业名称
- 联系地址/注册地址 -> 地址
- 联系人姓名 -> 联系人
- 法定代表人姓名 -> 法人姓名

只输出JSON对象"""

            system_prompt = "你是数据匹配工具，只输出JSON。"

            try:
                result = call_kimi_api(prompt, system_prompt, max_tokens=8192)
                result = re.sub(r'^```json\s*', '', result)
                result = re.sub(r'^```\s*', '', result)
                result = re.sub(r'```\s*$', '', result)
                json_match = re.search(r'\{[\s\S]*\}', result)
                if json_match:
                    result = json_match.group()
                parsed = json.loads(result)
                all_results.update(parsed)
            except Exception as e:
                print(f"  [WARN] 匹配失败: {e}")
                for tag in batch_tags:
                    all_results[tag] = "待补全"

    # 处理列表标签
    if list_tags:
        for list_tag in list_tags:
            # 提取列表名称
            match = re.search(r'for\s+item\s+in\s+(\S+)', list_tag)
            if match:
                list_name = match.group(1)
            else:
                list_name = list_tag.replace('列表', '')

            prompt = f"""【角色】企业信息匹配专家

【列表名称】：{list_name}

【企业信息】：
{company_info}

【任务】请从企业信息中提取所有符合条件的{list_name}数据，返回JSON数组。

例如项目负责人列表，请返回：
[
  {{"姓名": "xxx", "职务": "xxx", "职称": "xxx", ...}},
  {{"姓名": "yyy", "职务": "yyy", "职称": "yyy", ...}}
]

找不到返回空数组 []

只输出JSON数组"""

            system_prompt = "你是数据匹配工具，只输出JSON数组。"

            try:
                result = call_kimi_api(prompt, system_prompt, max_tokens=8192)
                result = re.sub(r'^```json\s*', '', result)
                result = re.sub(r'^```\s*', '', result)
                result = re.sub(r'```\s*$', '', result)

                # 尝试解析数组
                array_match = re.search(r'\[[\s\S]*\]', result)
                if array_match:
                    result = array_match.group()
                    parsed = json.loads(result)
                    all_results[list_name] = parsed
                else:
                    all_results[list_name] = []
            except Exception as e:
                print(f"  [WARN] 列表匹配失败: {e}")
                all_results[list_name] = []

    return all_results


def extract_tags_from_doc(doc: Document) -> list:
    """提取所有标签"""
    tags = set()
    pattern = r'\{\{(.*?)\}\}'

    for para in doc.paragraphs:
        matches = re.findall(pattern, para.text)
        tags.update(matches)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    matches = re.findall(pattern, para.text)
                    tags.update(matches)

    return sorted(list(tags))


def replace_text_in_paragraph(para, old_text: str, new_text) -> bool:
    if not old_text:
        return False
    new_text = str(new_text) if new_text is not None else ""
    full_text = para.text
    if old_text not in full_text:
        return False

    start_idx = full_text.find(old_text)
    end_idx = start_idx + len(old_text)

    run_start_idx = 0
    matching_runs = []

    for run in para.runs:
        run_end_idx = run_start_idx + len(run.text)
        if (run_start_idx <= start_idx < run_end_idx or
            run_start_idx < end_idx <= run_end_idx or
            (start_idx <= run_start_idx and run_end_idx <= end_idx)):
            matching_runs.append({
                'run': run,
                'run_start': run_start_idx,
                'run_end': run_end_idx,
                'text_start': max(start_idx, run_start_idx),
                'text_end': min(end_idx, run_end_idx)
            })
        run_start_idx = run_end_idx

    if not matching_runs:
        return False

    if len(matching_runs) == 1:
        run_info = matching_runs[0]
        run = run_info['run']
        run_local_start = run_info['text_start'] - run_info['run_start']
        run_local_end = run_info['text_end'] - run_info['run_start']
        run.text = run.text[:run_local_start] + new_text + run.text[run_local_end:]
        return True

    first_run_info = matching_runs[0]
    first_run = first_run_info['run']
    first_run_local_start = first_run_info['text_start'] - first_run_info['run_start']
    first_run.text = first_run.text[:first_run_local_start] + new_text

    for i, run_info in enumerate(matching_runs):
        if i == 0:
            continue
        run = run_info['run']
        run_local_end = run_info['text_end'] - run_info['run_start']
        if run_local_end < len(run.text):
            run.text = run.text[run_local_end:]
        else:
            run.text = ""
    return True


def apply_table_tags(table: Table, tagged_rows: list, unique_cells_info: list, is_dynamic: bool):
    """将打标结果应用回表格"""
    if is_dynamic:
        # 删除多余行，只保留表头和第一数据行
        rows_to_keep = len(tagged_rows)
        while len(table.rows) > rows_to_keep:
            tr = table.rows[-1]._tr
            tr.getparent().remove(tr)

    # 应用标签
    for row_idx, row in enumerate(table.rows):
        if row_idx >= len(tagged_rows):
            break

        tagged_cells = tagged_rows[row_idx]
        cells_info = unique_cells_info[row_idx] if row_idx < len(unique_cells_info) else []

        for cell_idx, cell in enumerate(row.cells):
            if cell_idx >= len(tagged_cells):
                break

            cell_info = cells_info[cell_idx] if cell_idx < len(cells_info) else {}

            if cell_info.get('is_merged'):
                continue

            tagged_text = tagged_cells[cell_idx]

            # 清空单元格并写入新内容
            for para in cell.paragraphs:
                for run in para.runs:
                    run.text = ""

            if cell.paragraphs:
                if cell.paragraphs[0].runs:
                    cell.paragraphs[0].runs[0].text = tagged_text
                else:
                    cell.paragraphs[0].add_run(tagged_text)


# ============== 主流程 ==============

print("\n【Step 1】加载文档")
print("-" * 70)

doc = Document(TEMPLATE_PATH)
print(f"[OK] 文档加载成功")
print(f"  - 段落数量: {len(doc.paragraphs)}")
print(f"  - 表格数量: {len(doc.tables)}")

with open(COMPANY_INFO_PATH, 'r', encoding='utf-8') as f:
    company_info = f.read()
print(f"[OK] 企业信息已加载，共 {len(company_info)} 字符")

print("\n【Step 2】合并碎片化Run")
print("-" * 70)

merged = merge_adjacent_runs(doc)
print(f"[OK] Run合并完成，共合并 {merged} 个")

print("\n【Step 3】AI辅助打标 - 段落")
print("-" * 70)

valid_paragraphs = [(i, p) for i, p in enumerate(doc.paragraphs) if p.text.strip()]
print(f"发现 {len(valid_paragraphs)} 个有效段落")

batch_size = 20
paragraph_processed = 0

for batch_start in range(0, len(valid_paragraphs), batch_size):
    batch_paras = valid_paragraphs[batch_start:batch_start + batch_size]
    print(f"  处理批次 {batch_start//batch_size + 1}/{(len(valid_paragraphs)+batch_size-1)//batch_size}...")

    try:
        tagged_map = call_kimi_tag_paragraphs_batch(batch_paras)

        for orig_idx, para in batch_paras:
            if orig_idx in tagged_map:
                new_text = tagged_map[orig_idx]
                old_text = para.text.strip()

                if "{{" in new_text and new_text != old_text:
                    # 替换整个段落内容
                    full_text = para.text
                    replace_text_in_paragraph(para, full_text, new_text)
                    paragraph_processed += 1

        print(f"    本批次打标 {len(tagged_map)} 个")
        time.sleep(1)

    except Exception as e:
        print(f"    [WARN] 批次处理失败: {e}")

print(f"[OK] 段落打标完成，成功处理 {paragraph_processed} 个")

print("\n【Step 4】AI辅助打标 - 表格（生成docxtpl循环标签）")
print("-" * 70)

tables_with_context = get_tables_with_upper_context(doc)
print(f"发现 {len(tables_with_context)} 个表格")

table_processed = 0
list_info = {}  # 存储列表信息

for table_idx, table_ctx in enumerate(tables_with_context):
    table = table_ctx['table']
    upper_context = table_ctx['upper_context']
    serialized_table = table_ctx['serialized_table']
    unique_cells_info = table_ctx['unique_cells_info']

    is_dynamic = detect_dynamic_list(table_ctx)
    print(f"  处理表格 {table_idx + 1} (动态列表: {is_dynamic})...")

    try:
        tagged_table = call_kimi_tag_table(serialized_table, upper_context, is_dynamic)
        tagged_rows = parse_serialized_table(tagged_table)

        apply_table_tags(table, tagged_rows, unique_cells_info, is_dynamic)

        # 记录列表信息
        if is_dynamic:
            for row in tagged_rows:
                for cell in row:
                    if 'tr for' in cell:
                        match = re.search(r'tr\s+for\s+item\s+in\s+(\S+)', cell)
                        if match:
                            list_name = match.group(1)
                            list_info[list_name] = True

        table_processed += 1
        print(f"    完成")
        time.sleep(1)

    except Exception as e:
        print(f"    [WARN] 失败: {e}")

print(f"[OK] 表格打标完成，成功处理 {table_processed} 个")
print(f"发现动态列表: {list(list_info.keys())}")

print("\n【Step 5】标签提取")
print("-" * 70)

tags = extract_tags_from_doc(doc)
print(f"[OK] 提取到 {len(tags)} 个标签")
print(f"  标签列表: {tags}")

print("\n【Step 6】企业信息匹配")
print("-" * 70)

start_time = time.time()
matched_data = call_kimi_match_info(tags, company_info)
elapsed = time.time() - start_time
print(f"[OK] 匹配完成，耗时 {elapsed:.2f}秒")

matched_count = sum(1 for v in matched_data.values() if v not in ["待补全", []] and v)
print(f"  成功匹配: {matched_count}/{len(tags)} 个")

print("\n【Step 7】构建docxtpl上下文")
print("-" * 70)

context = {}
for tag, value in matched_data.items():
    context[tag] = value

print(f"上下文构建完成，共 {len(context)} 个字段")

print("\n【Step 8】使用docxtpl渲染动态表格")
print("-" * 70)

# 保存带标签的模板
temp_template_path = r"D:\桌面\temp_template.docx"
doc.save(temp_template_path)
print(f"  临时模板已保存: {temp_template_path}")

# 使用docxtpl渲染
try:
    tpl = DocxTemplate(temp_template_path)
    tpl.render(context)
    tpl.save(OUTPUT_PATH)
    print(f"[OK] docxtpl渲染完成")
except Exception as e:
    print(f"[WARN] docxtpl渲染失败，使用普通方式保存: {e}")
    doc.save(OUTPUT_PATH)

# 清理临时文件
try:
    os.remove(temp_template_path)
except:
    pass

print("\n" + "=" * 70)
print("完整流程处理完成!")
print("=" * 70)
print(f"结束时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
print(f"\n处理结果摘要:")
print(f"  - 段落打标: {paragraph_processed} 个")
print(f"  - 表格打标: {table_processed} 个")
print(f"  - 标签提取: {len(tags)} 个")
print(f"  - 信息匹配: {matched_count}/{len(tags)} 个成功")
print(f"  - 动态列表: {len(list_info)} 个")
print(f"\n输出文件: {OUTPUT_PATH}")
