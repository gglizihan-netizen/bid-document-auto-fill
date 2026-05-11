# -*- coding: utf-8 -*-
"""
验证输出文档的填充效果
"""

import os
import sys
import re
from pathlib import Path

if sys.platform == 'win32':
    import io
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from docx import Document

OUTPUT_PATH = r"D:\桌面\最终输出_完整填充_v4.docx"
TEMPLATE_PATH = r"D:\桌面\投标文件组成模版01.docx"

print("=" * 70)
print("输出文档填充效果验证")
print("=" * 70)

# 加载输出文档
doc = Document(OUTPUT_PATH)

print("\n【段落内容检查】")
print("-" * 70)

# 检查是否还有未填充的标签
remaining_tags = set()
pattern = r'\{\{(.*?)\}\}'

for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if not text:
        continue

    # 检查是否有剩余标签
    matches = re.findall(pattern, text)
    if matches:
        remaining_tags.update(matches)
        print(f"[段落 {i}] 还有标签: {matches}")
        print(f"  内容: {text[:80]}...")

    # 检查是否还有明显空白处
    if re.search(r'_{2,}', text):
        print(f"[段落 {i}] 还有下划线空白")
        print(f"  内容: {text[:80]}...")

    if re.search(r'\s{4,}', text):
        # 排除合理的空格
        if not re.search(r'年\s+月\s+日', text):
            print(f"[段落 {i}] 还有大空白")
            print(f"  内容: {text[:80]}...")

print(f"\n剩余未填充标签数量: {len(remaining_tags)}")
if remaining_tags:
    print(f"剩余标签: {sorted(remaining_tags)}")

print("\n【表格内容检查】")
print("-" * 70)

def is_blank(text):
    if not text or not text.strip():
        return True
    text = text.strip()
    if text in ['[空]', '待填写', '待补全', '[合并]']:
        return True
    if re.match(r'^_{2,}$', text):
        return True
    return False

for table_idx, table in enumerate(doc.tables):
    print(f"\n表格 {table_idx + 1}:")

    # 统计单元格状态
    total_cells = 0
    blank_cells = 0
    filled_cells = 0
    tag_cells = 0

    for row_idx, row in enumerate(table.rows[:5]):  # 只显示前5行
        row_status = []
        for cell in row.cells:
            total_cells += 1
            text = cell.text.strip()

            if is_blank(text):
                blank_cells += 1
                row_status.append('[空]')
            elif '{{' in text and '}}' in text:
                tag_cells += 1
                row_status.append('[标签]')
            else:
                filled_cells += 1
                preview = text[:15] + '...' if len(text) > 15 else text
                row_status.append(preview)

        print(f"  行{row_idx}: {row_status}")

    # 统计剩余表格
    remaining_rows = len(table.rows) - 5
    if remaining_rows > 0:
        for row in table.rows[5:]:
            for cell in row.cells:
                total_cells += 1
                text = cell.text.strip()
                if is_blank(text):
                    blank_cells += 1
                elif '{{' in text and '}}' in text:
                    tag_cells += 1
                else:
                    filled_cells += 1

    print(f"  统计: 总{total_cells} 空{blank_cells} 标签{tag_cells} 已填{filled_cells}")

print("\n【填充效果评估】")
print("-" * 70)

# 对比原始模板
template_doc = Document(TEMPLATE_PATH)
template_blank_count = 0
output_filled_count = 0

# 统计模板空白
for para in template_doc.paragraphs:
    if re.search(r'_{2,}|\s{4,}', para.text):
        template_blank_count += 1

# 统计输出填充
for para in doc.paragraphs:
    text = para.text
    # 检查是否有实际内容填充（不是标签也不是空白）
    if text.strip() and not re.search(r'\{\{.*?\}\}|_{2,}|\s{4,}', text):
        output_filled_count += 1

print(f"模板空白段落估计: {template_blank_count}")
print(f"输出文档有内容的段落: {output_filled_count}")

print("\n" + "=" * 70)
print("验证完成")
print("=" * 70)

if remaining_tags:
    print(f"\n[注意] 还有 {len(remaining_tags)} 个标签未填充")
    print("请检查企业信息是否包含相关内容")
else:
    print("\n[OK] 所有标签都已填充")