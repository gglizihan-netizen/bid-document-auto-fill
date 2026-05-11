# -*- coding: utf-8 -*-
"""
完整功能测试 - 增强版
"""

import os
import sys
import re
import json
import time
from pathlib import Path

if sys.platform == 'win32':
    import io
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8')

from docx import Document
from openai import OpenAI
from dotenv import load_dotenv

load_dotenv()

KIMI_API_KEY = os.getenv('KIMI_API_KEY', '')
KIMI_BASE_URL = os.getenv('KIMI_BASE_URL', 'https://api.moonshot.cn/v1')
KIMI_MODEL = 'kimi-k2.6'

TEMPLATE_PATH = r'E:\Users\han\Desktop\投标文件组成01.docx'
COMPANY_INFO_PATH = r'E:\Users\han\Desktop\企业投标信息01.txt'
OUTPUT_PATH = r'E:\Users\han\Desktop\auto_fill\outputs\test_output_v3.docx'

print("=" * 60)
print("完整功能测试 - 增强版")
print(f"使用模型: {KIMI_MODEL}")
print("=" * 60)

client = OpenAI(api_key=KIMI_API_KEY, base_url=KIMI_BASE_URL)

# 加载文档
doc = Document(TEMPLATE_PATH)
print(f"文档: {len(doc.paragraphs)} 段落, {len(doc.tables)} 表格")

with open(COMPANY_INFO_PATH, 'r', encoding='utf-8') as f:
    company_info = f.read()
print(f"企业信息: {len(company_info)} 字符")

# Step 1: 检测并标记空白处
print()
print("【Step 1】检测空白处")
print("-" * 60)

def has_blank(text):
    """检测文本中是否有空白处"""
    # 下划线
    if re.search(r'_{2,}', text):
        return True
    # 括号提示（排除注释性括号）
    if re.search(r'（[^）]{0,10}）', text):
        # 排除"注："开头的注释
        if not text.startswith('（注：'):
            return True
    # 冒号后空白
    if re.search(r'：\s{3,}', text):
        return True
    # 日期空格
    if re.search(r'\s{2,}年\s{2,}月', text):
        return True
    # 空格填充
    if re.search(r'\s{5,}', text):
        return True
    return False

valid_paragraphs = []
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if text and has_blank(text):
        valid_paragraphs.append((i, para))

print(f"检测到 {len(valid_paragraphs)} 个包含空白处的段落")

# Step 2: AI打标
print()
print("【Step 2】AI打标")
print("-" * 60)

all_tags = set()
paragraph_processed = 0

for batch_start in range(0, len(valid_paragraphs), 5):
    batch_paras = valid_paragraphs[batch_start:batch_start + 5]
    text_block = "\n".join([f"[{idx}] {p.text.strip()}" for idx, p in batch_paras])

    prompt = f"""你是投标文件处理专家。请识别下面文本中所有需要填写的空白处，并用具体标签替换。

【空白处类型】
1. 下划线:________
2. 括号提示:(项目名称)、（单位盖章）
3. 冒号后空白:投标人:____
4. 日期空格:  年  月  日

【打标规则】
1. 用{{{{标签名}}}}替换空白处
2. 标签必须具体明确:
   - 错误: {{{{姓名}}}}、{{{{地址}}}}
   - 正确: {{{{法定代表人姓名}}}}、{{{{投标人地址}}}}
3. 保留原文其他内容不变
4. 保留每行开头的[数字]索引

【示例】
输入: [1] 投标人：        （单位盖章）
输出: [1] 投标人：{{{{投标人单位名称}}}}（单位盖章）

【待处理文本】
{text_block}"""

    try:
        response = client.chat.completions.create(
            model=KIMI_MODEL,
            messages=[
                {"role": "system", "content": "你是投标文件处理专家，负责识别空白处并打标签。"},
                {"role": "user", "content": prompt}
            ],
            max_tokens=2000
        )
        result = response.choices[0].message.content

        # 解析结果
        pattern = r"^\[(\d+)\]\s*(.*)$"
        for line in result.splitlines():
            m = re.match(pattern, line.strip())
            if m:
                idx = int(m.group(1))
                content = m.group(2)
                for orig_idx, para in batch_paras:
                    if orig_idx == idx and "{{" in content:
                        para.clear()
                        para.add_run(content)
                        paragraph_processed += 1
                        tags = re.findall(r"\{\{([^}]+)\}\}", content)
                        all_tags.update(tags)

        print(f"  批次 {batch_start//5 + 1}: OK, 标签数: {len(all_tags)}")
        time.sleep(0.5)

    except Exception as e:
        print(f"  批次 {batch_start//5 + 1} 失败: {e}")

print(f"段落打标完成: {paragraph_processed} 个段落")
print(f"提取到标签: {len(all_tags)} 个")
print(f"标签列表: {sorted(list(all_tags))}")

# Step 3: 表格打标
print()
print("【Step 3】表格打标")
print("-" * 60)

table_tags = set()

for table_idx, table in enumerate(doc.tables):
    # 序列化表格
    rows_text = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            text = cell.text.strip()
            if not text or text == '':
                cells.append('[空]')
            else:
                cells.append(text.replace('|', '｜')[:30])
        rows_text.append('| ' + ' | '.join(cells) + ' |')

    serialized = '\n'.join(rows_text[:5])  # 只取前5行

    # 跳过没有空单元格的表格
    if '[空]' not in serialized:
        print(f"  表格 {table_idx}: 无空单元格，跳过")
        continue

    prompt = f"""你是表格打标专家。请为下面的表格空白处打标签。

【表格内容】
{serialized}

【打标规则】
1. 用{{{{具体标签名}}}}替换[空]
2. 标签必须结合表格上下文，如:
   - 项目负责人表格: {{{{项目负责人姓名}}}}
   - 投标人信息表格: {{{{投标人联系电话}}}}
3. 保留表格结构

【输出】
只输出打标后的表格文本"""

    try:
        response = client.chat.completions.create(
            model=KIMI_MODEL,
            messages=[
                {"role": "system", "content": "你是表格打标专家。"},
                {"role": "user", "content": prompt}
            ],
            max_tokens=2000
        )
        result = response.choices[0].message.content

        # 提取标签
        tags = re.findall(r"\{\{([^}]+)\}\}", result)
        table_tags.update(tags)

        # 应用到表格
        tagged_rows = []
        for line in result.split('\n'):
            if '|' in line:
                cells = [c.strip() for c in line.split('|') if c.strip()]
                tagged_rows.append(cells)

        # 写回表格
        for row_idx, row in enumerate(table.rows):
            if row_idx >= len(tagged_rows):
                break
            tagged_cells = tagged_rows[row_idx]
            for cell_idx, cell in enumerate(row.cells):
                if cell_idx >= len(tagged_cells):
                    break
                tagged_text = tagged_cells[cell_idx]
                if '{{' in tagged_text:
                    for para in cell.paragraphs:
                        if para.text.strip() == '' or para.text.strip() == '[空]':
                            para.clear()
                            para.add_run(tagged_text)

        print(f"  表格 {table_idx}: OK, 标签: {tags[:5]}")
        time.sleep(0.5)

    except Exception as e:
        print(f"  表格 {table_idx} 失败: {e}")

all_tags.update(table_tags)
print(f"表格打标完成，新增标签: {len(table_tags)} 个")

# Step 4: 信息匹配
print()
print("【Step 4】信息匹配")
print("-" * 60)

tags_list = sorted(list(all_tags))
matched_data = {}

if tags_list:
    prompt = f"""你是企业信息匹配专家。请为每个标签从企业信息中找到对应值。

【标签列表】
{json.dumps(tags_list, ensure_ascii=False)}

【企业信息】
{company_info}

【匹配规则】
1. 精确匹配标签语义
2. 日期类标签提取数字
3. 找不到返回"待补全"

【输出格式】
JSON对象"""

    try:
        response = client.chat.completions.create(
            model=KIMI_MODEL,
            messages=[
                {"role": "system", "content": "你是企业信息匹配专家。"},
                {"role": "user", "content": prompt}
            ],
            max_tokens=4000
        )
        result = response.choices[0].message.content

        result = re.sub(r"^```json\s*", "", result)
        result = re.sub(r"^```\s*", "", result)
        result = re.sub(r"```\s*$", "", result)

        json_match = re.search(r"\{[\s\S]*\}", result)
        if json_match:
            matched_data = json.loads(json_match.group())
            for tag in tags_list:
                if tag not in matched_data:
                    matched_data[tag] = "待补全"
    except Exception as e:
        print(f"匹配失败: {e}")
        matched_data = {tag: "待补全" for tag in tags_list}

filled = sum(1 for v in matched_data.values() if v and v != "待补全")
total = len(matched_data)
print(f"匹配结果: {filled}/{total} 成功 ({filled*100//total if total else 0}%)")

print()
print("匹配结果预览:")
for i, (k, v) in enumerate(list(matched_data.items())[:15]):
    v_preview = str(v)[:40] if isinstance(v, str) else str(v)[:60]
    print(f"  {k}: {v_preview}")

# Step 5: 填充文档
print()
print("【Step 5】填充文档")
print("-" * 60)

fill_count = 0

# 填充段落
for para in doc.paragraphs:
    for tag, value in matched_data.items():
        placeholder = f"{{{{{tag}}}}}"
        if placeholder in para.text:
            full_text = para.text
            new_text = full_text.replace(placeholder, str(value))
            if new_text != full_text:
                para.clear()
                para.add_run(new_text)
                fill_count += 1

# 填充表格
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for tag, value in matched_data.items():
                    placeholder = f"{{{{{tag}}}}}"
                    if placeholder in para.text:
                        full_text = para.text
                        new_text = full_text.replace(placeholder, str(value))
                        if new_text != full_text:
                            para.clear()
                            para.add_run(new_text)
                            fill_count += 1

print(f"填充完成: {fill_count} 处")

# Step 6: 保存
print()
print("【Step 6】保存结果")
print("-" * 60)

Path("outputs").mkdir(exist_ok=True)
doc.save(OUTPUT_PATH)
print(f"输出保存到: {OUTPUT_PATH}")

# 最终统计
print()
print("=" * 60)
print("测试结果总结")
print("=" * 60)
print(f"段落打标: {paragraph_processed} 个")
print(f"表格打标: {len(table_tags)} 个标签")
print(f"标签总数: {len(all_tags)} 个")
print(f"信息匹配: {filled}/{total} 成功 ({filled*100//total if total else 0}%)")
print(f"文档填充: {fill_count} 处")
