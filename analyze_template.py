# -*- coding: utf-8 -*-
"""
分析模板文档中所有需要填充的位置
"""

import os
import sys
import re
from pathlib import Path

if sys.platform == 'win32':
    import io
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from docx import Document

TEMPLATE_PATH = r"D:\桌面\投标文件组成模版01.docx"

doc = Document(TEMPLATE_PATH)

print("=" * 80)
print("模板文档详细分析")
print("=" * 80)

# 分析段落
print("\n" + "=" * 80)
print("【段落分析】")
print("=" * 80)

blank_patterns = [
    r'_{2,}',           # 下划线
    r'\s{3,}',          # 多个空格
    r'（[^）]*）',      # 中文括号提示
    r'\([^)]*\)',       # 英文括号提示
    r'\[[^\]]*\]',      # 方括号提示
    r'：\s*$',          # 冒号后留白
    r':\s*$',           # 英文冒号后留白
]

paragraphs_with_blanks = []
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if not text:
        continue

    # 检查是否包含空白处
    has_blank = False
    blank_types = []

    for pattern in blank_patterns:
        matches = re.findall(pattern, text)
        if matches:
            has_blank = True
            blank_types.append(f"{pattern}: {matches}")

    if has_blank:
        paragraphs_with_blanks.append({
            'index': i,
            'text': text,
            'blank_types': blank_types
        })

print(f"\n发现 {len(paragraphs_with_blanks)} 个可能包含空白的段落：\n")
for item in paragraphs_with_blanks:
    print(f"[段落 {item['index']}]")
    print(f"  文本: {item['text'][:80]}...")
    print(f"  空白类型: {item['blank_types'][:2]}")
    print()

# 分析表格
print("=" * 80)
print("【表格分析】")
print("=" * 80)

def is_blank(text):
    if not text or not text.strip():
        return True
    text = text.strip()
    if text in ['[空]', '待填写', '待补全']:
        return True
    if re.match(r'^_{2,}$', text):
        return True
    if re.match(r'^\s*$', text):
        return True
    return False

total_cells = 0
blank_cells = 0

for table_idx, table in enumerate(doc.tables):
    print(f"\n【表格 {table_idx + 1}】")

    table_total = 0
    table_blank = 0
    blank_positions = []

    for row_idx, row in enumerate(table.rows):
        for cell_idx, cell in enumerate(row.cells):
            table_total += 1
            total_cells += 1

            text = cell.text.strip()
            if is_blank(text):
                table_blank += 1
                blank_cells += 1
                blank_positions.append(f"行{row_idx}列{cell_idx}")

    print(f"  总单元格: {table_total}, 空白单元格: {table_blank}")
    print(f"  空白比例: {table_blank/table_total*100:.1f}%")
    if blank_positions[:10]:
        print(f"  空白位置: {blank_positions[:10]}...")

print("\n" + "=" * 80)
print("【汇总统计】")
print("=" * 80)
print(f"总段落数: {len(doc.paragraphs)}")
print(f"包含空白的段落: {len(paragraphs_with_blanks)}")
print(f"总表格数: {len(doc.tables)}")
print(f"总单元格数: {total_cells}")
print(f"空白单元格数: {blank_cells}")
print(f"空白单元格比例: {blank_cells/total_cells*100:.1f}%")
