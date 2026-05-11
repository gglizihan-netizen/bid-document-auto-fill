# -*- coding: utf-8 -*-
"""
资信标自动填充系统 - 完整流程测试
端到端测试整个填充流程
"""

import os
import sys
import re
import json
import time
from pathlib import Path
from datetime import datetime

# 设置控制台编码
if sys.platform == 'win32':
    import io
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8')

# 添加项目路径
sys.path.insert(0, str(Path(__file__).parent))

from docx import Document
from dotenv import load_dotenv
from openai import OpenAI

# 加载环境变量
load_dotenv()

# 配置
KIMI_API_KEY = os.getenv("KIMI_API_KEY", "")
KIMI_BASE_URL = os.getenv("KIMI_BASE_URL", "https://api.moonshot.cn/v1")
KIMI_MODEL = os.getenv("KIMI_MODEL", "moonshot-v1-8k")

# 测试文件路径
TEMPLATE_PATH = r"D:\桌面\投标文件组成模版01.docx"
COMPANY_INFO_PATH = r"D:\桌面\企业信息01.txt"
OUTPUT_PATH = r"D:\桌面\测试输出_完整填充结果.docx"

print("=" * 60)
print("资信标自动填充系统 - 完整流程测试")
print("=" * 60)
print(f"开始时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

# 初始化客户端
client = OpenAI(api_key=KIMI_API_KEY, base_url=KIMI_BASE_URL)

# ============== 核心函数 ==============

def get_run_format(run):
    font = run.font
    return (
        font.name,
        font.size.pt if font.size else None,
        font.bold,
        font.italic,
        font.underline,
    )

def merge_adjacent_runs(doc: Document) -> int:
    """合并格式相同的相邻Run"""
    merged_count = 0

    for para in doc.paragraphs:
        if len(para.runs) <= 1:
            continue
        i = 0
        while i < len(para.runs) - 1:
            current_run = para.runs[i]
            next_run = para.runs[i + 1]
            if get_run_format(current_run) == get_run_format(next_run):
                current_run.text += next_run.text
                next_run._element.getparent().remove(next_run._element)
                merged_count += 1
            else:
                i += 1

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    if len(para.runs) <= 1:
                        continue
                    i = 0
                    while i < len(para.runs) - 1:
                        current_run = para.runs[i]
                        next_run = para.runs[i + 1]
                        if get_run_format(current_run) == get_run_format(next_run):
                            current_run.text += next_run.text
                            next_run._element.getparent().remove(next_run._element)
                            merged_count += 1
                        else:
                            i += 1

    return merged_count

def replace_text_in_paragraph(para, old_text: str, new_text) -> bool:
    """跨Run替换文本，保留格式"""
    if not old_text:
        return False

    new_text = str(new_text) if new_text is not None else ""
    full_text = para.text

    if old_text not in full_text:
        return False

    start_idx = full_text.find(old_text)
    end_idx = start_idx + len(old_text)

    run_start_idx = 0
    matching_runs = []

    for run in para.runs:
        run_end_idx = run_start_idx + len(run.text)
        if (run_start_idx <= start_idx < run_end_idx or
            run_start_idx < end_idx <= run_end_idx or
            (start_idx <= run_start_idx and run_end_idx <= end_idx)):
            matching_runs.append({
                'run': run,
                'run_start': run_start_idx,
                'run_end': run_end_idx,
                'text_start': max(start_idx, run_start_idx),
                'text_end': min(end_idx, run_end_idx)
            })
        run_start_idx = run_end_idx

    if not matching_runs:
        return False

    if len(matching_runs) == 1:
        run_info = matching_runs[0]
        run = run_info['run']
        run_local_start = run_info['text_start'] - run_info['run_start']
        run_local_end = run_info['text_end'] - run_info['run_start']
        run.text = run.text[:run_local_start] + new_text + run.text[run_local_end:]
        return True

    first_run_info = matching_runs[0]
    first_run = first_run_info['run']
    first_run_local_start = first_run_info['text_start'] - first_run_info['run_start']

    new_run_text = first_run.text[:first_run_local_start] + new_text
    first_run.text = new_run_text

    for i, run_info in enumerate(matching_runs):
        if i == 0:
            continue
        run = run_info['run']
        run_local_start = run_info['text_start'] - run_info['run_start']
        run_local_end = run_info['text_end'] - run_info['run_start']

        if run_local_end < len(run.text):
            run.text = run.text[run_local_end:]
        else:
            run.text = ""

    return True

def call_kimi_tagging(text: str, is_table: bool = False) -> str:
    """调用Kimi API进行打标"""
    if is_table:
        prompt = f"""【角色】你是专业的投标文件标签生成器。

【任务】为Markdown表格中的空白单元格打标签。

【输入表格】：
{text}

【打标规则】
1. 只在 [空] 或空白单元格处填入 {{{{标签名}}}}
2. 标签名根据该列表头语义生成，简洁明确
3. 如果同一列有多个空白，使用相同标签
4. 保持表格格式完整

【重要】只输出打标后的Markdown表格，不要输出任何说明文字！"""
    else:
        prompt = f"""【角色】你是专业的投标文件标签生成器。

【任务】识别文本中的空白处并打上标签。

【输入文本】：
{text}

【空白处识别】
1. 下划线：________、____、_等
2. 多个空格：      （连续空格）
3. 冒号后留白：投标人：  （冒号后无内容）
4. 括号提示：(项目名称)、[请填写]、（单位盖章）等
5. 日期格式空白：  年  月  日

【打标规则】
1. 根据上下文语义生成简洁的中文标签名
2. 标签格式：{{{{标签名}}}}
3. 必须原样返回整句，只替换空白处为标签
4. 严禁删减、增加或改动任何原有字符

【重要】只输出打标后的文本，不要输出任何说明！"""

    try:
        response = client.chat.completions.create(
            model=KIMI_MODEL,
            messages=[
                {"role": "system", "content": "你是标签生成器，只输出打标后的内容。"},
                {"role": "user", "content": prompt}
            ],
            temperature=0.1,
            max_tokens=2048
        )
        result = response.choices[0].message.content.strip()

        # 清理AI返回的多余内容
        result = re.sub(r'^```markdown\s*', '', result)
        result = re.sub(r'^```\s*', '', result)
        result = re.sub(r'```\s*$', '', result)

        return result.strip()
    except Exception as e:
        print(f"  [WARN] API调用失败: {e}")
        return text

def call_kimi_match_info(tags: list, company_info: str) -> dict:
    """调用Kimi API进行信息匹配"""
    if not tags:
        return {}

    prompt = f"""【角色】你是一名专业的企业信息匹配专家。

【标签列表】（共{len(tags)}个）：
{json.dumps(tags, ensure_ascii=False, indent=2)}

【企业信息】：
{company_info[:3000]}

【输出要求】：
只输出一个JSON对象，格式如下：
{{"标签1": "对应值1", "标签2": "对应值2", ...}}

找不到对应信息的标签，value填"待补全"。"""

    try:
        response = client.chat.completions.create(
            model=KIMI_MODEL,
            messages=[
                {"role": "system", "content": "你是数据匹配工具，只输出JSON。"},
                {"role": "user", "content": prompt}
            ],
            temperature=0.1,
            max_tokens=4096
        )

        content = response.choices[0].message.content.strip()

        # 清理可能的markdown代码块
        content = re.sub(r'^```json\s*', '', content)
        content = re.sub(r'^```\s*', '', content)
        content = re.sub(r'```\s*$', '', content)

        json_match = re.search(r'\{[\s\S]*\}', content)
        if json_match:
            content = json_match.group()

        return json.loads(content)
    except Exception as e:
        print(f"  [WARN] 信息匹配失败: {e}")
        return {tag: "待补全" for tag in tags}

def extract_tags_from_doc(doc: Document) -> list:
    """从文档中提取所有{{标签}}"""
    tags = set()
    pattern = r'\{\{(.*?)\}\}'

    for para in doc.paragraphs:
        matches = re.findall(pattern, para.text)
        tags.update(matches)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    matches = re.findall(pattern, para.text)
                    tags.update(matches)

    return sorted(list(tags))

def table_to_markdown(table) -> str:
    """将Word表格转换为Markdown格式"""
    rows = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            text = cell.text.strip()
            if not text or text in ['待补全', '待填写', '[空]']:
                cells.append('[空]')
            else:
                text = text.replace('|', '｜').replace('\n', ' ')
                cells.append(text)
        rows.append('| ' + ' | '.join(cells) + ' |')

    if rows:
        col_count = len(table.rows[0].cells) if table.rows else 0
        separator = '|' + '|'.join(['---'] * col_count) + '|'
        rows.insert(1, separator)

    return '\n'.join(rows)

def parse_markdown_table(md_text: str) -> list:
    """解析Markdown表格为二维数组"""
    lines = [line.strip() for line in md_text.strip().split('\n') if line.strip()]
    rows = []
    for line in lines:
        if '|' in line and '---' in line:
            continue
        cells = [cell.strip() for cell in line.split('|')]
        cells = [cell for cell in cells if cell]
        if cells:
            rows.append(cells)
    return rows

def apply_markdown_tags_to_table(table, md_rows: list):
    """将Markdown打标结果应用回Word表格"""
    for row_idx, row in enumerate(table.rows):
        if row_idx >= len(md_rows):
            break

        md_cells = md_rows[row_idx]
        for cell_idx, cell in enumerate(row.cells):
            if cell_idx >= len(md_cells):
                break

            md_text = md_cells[cell_idx]

            if '{{' in md_text and '}}' in md_text:
                for para in cell.paragraphs:
                    if para.text.strip():
                        replace_text_in_paragraph(para, para.text, md_text)
                    else:
                        if para.runs:
                            para.runs[0].text = md_text
                        else:
                            para.add_run(md_text)

# ============== 主流程 ==============

print("\n【Step 1】加载文档和配置")
print("-" * 40)

# 加载文档
doc = Document(TEMPLATE_PATH)
print(f"[OK] 文档加载成功")
print(f"  - 段落数量: {len(doc.paragraphs)}")
print(f"  - 表格数量: {len(doc.tables)}")

# 读取企业信息
with open(COMPANY_INFO_PATH, 'r', encoding='utf-8') as f:
    company_info = f.read()
print(f"[OK] 企业信息已加载，共 {len(company_info)} 字符")

print("\n【Step 2】合并碎片化Run")
print("-" * 40)
start_time = time.time()
merged = merge_adjacent_runs(doc)
elapsed = time.time() - start_time
print(f"[OK] Run合并完成，共合并 {merged} 个，耗时 {elapsed:.2f}秒")

print("\n【Step 3】AI辅助打标 - 段落")
print("-" * 40)
start_time = time.time()

# 处理正文段落
paragraphs_to_process = []
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if text and len(text) < 500:
        # 检查是否可能包含空白
        if any(c in text for c in ['_', '：', ':', '[', '□', '空', '待填写', '待补全']):
            paragraphs_to_process.append((i, para))

print(f"发现 {len(paragraphs_to_process)} 个可能包含空白的段落")

max_paragraphs = 15  # 限制处理数量以节省API调用
processed_count = 0

for idx, (para_idx, para) in enumerate(paragraphs_to_process[:max_paragraphs]):
    try:
        original_text = para.text
        tagged_text = call_kimi_tagging(original_text, is_table=False)

        if "{{" in tagged_text and tagged_text != original_text:
            if replace_text_in_paragraph(para, original_text, tagged_text):
                processed_count += 1
                print(f"  [{idx+1}/{min(len(paragraphs_to_process), max_paragraphs)}] 段落{para_idx}: 打标成功")

        if (idx + 1) % 5 == 0:
            print(f"  进度: {idx + 1}/{min(len(paragraphs_to_process), max_paragraphs)}")

    except Exception as e:
        print(f"  [WARN] 打标段落 {para_idx} 失败: {e}")

elapsed = time.time() - start_time
print(f"[OK] 段落打标完成，成功处理 {processed_count} 个，耗时 {elapsed:.2f}秒")

print("\n【Step 4】AI辅助打标 - 表格")
print("-" * 40)
start_time = time.time()

table_processed = 0
for table_idx, table in enumerate(doc.tables[:5]):  # 只处理前5个表格
    try:
        # 转换为Markdown
        md_table = table_to_markdown(table)

        # 调用AI打标
        tagged_md = call_kimi_tagging(md_table, is_table=True)

        # 解析Markdown结果
        md_rows = parse_markdown_table(tagged_md)

        # 应用回Word表格
        apply_markdown_tags_to_table(table, md_rows)

        table_processed += 1
        print(f"  表格{table_idx + 1}: 处理完成")

    except Exception as e:
        print(f"  [WARN] 处理表格 {table_idx} 失败: {e}")

elapsed = time.time() - start_time
print(f"[OK] 表格打标完成，成功处理 {table_processed} 个，耗时 {elapsed:.2f}秒")

print("\n【Step 5】标签提取")
print("-" * 40)
start_time = time.time()
tags = extract_tags_from_doc(doc)
elapsed = time.time() - start_time
print(f"[OK] 提取到 {len(tags)} 个标签")
print(f"  标签列表: {tags[:20]}{'...' if len(tags) > 20 else ''}")
print(f"  耗时: {elapsed:.2f}秒")

print("\n【Step 6】企业信息匹配")
print("-" * 40)
start_time = time.time()
matched_data = call_kimi_match_info(tags, company_info)
elapsed = time.time() - start_time
print(f"[OK] 信息匹配完成，耗时 {elapsed:.2f}秒")

# 显示匹配结果统计
matched_count = sum(1 for v in matched_data.values() if v != "待补全")
print(f"  成功匹配: {matched_count}/{len(tags)} 个标签")

# 显示部分匹配结果
print("\n  匹配结果预览:")
for i, (tag, value) in enumerate(list(matched_data.items())[:10]):
    print(f"    {tag}: {value}")

print("\n【Step 7】内容填充")
print("-" * 40)
start_time = time.time()
replacement_count = 0

# 处理正文
for para_idx, para in enumerate(doc.paragraphs):
    for tag, value in matched_data.items():
        placeholder = f"{{{{{tag}}}}}"
        if placeholder in para.text:
            if replace_text_in_paragraph(para, placeholder, value):
                replacement_count += 1

# 处理表格
for table_idx, table in enumerate(doc.tables):
    for row_idx, row in enumerate(table.rows):
        for cell_idx, cell in enumerate(row.cells):
            for para in cell.paragraphs:
                for tag, value in matched_data.items():
                    placeholder = f"{{{{{tag}}}}}"
                    if placeholder in para.text:
                        if replace_text_in_paragraph(para, placeholder, value):
                            replacement_count += 1

elapsed = time.time() - start_time
print(f"[OK] 内容填充完成，共替换 {replacement_count} 处，耗时 {elapsed:.2f}秒")

print("\n【Step 8】保存结果")
print("-" * 40)
doc.save(OUTPUT_PATH)
print(f"[OK] 文档已保存至: {OUTPUT_PATH}")

print("\n" + "=" * 60)
print("完整流程测试完成!")
print("=" * 60)
print(f"结束时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
print(f"\n测试结果摘要:")
print(f"  - 段落打标: {processed_count} 个")
print(f"  - 表格打标: {table_processed} 个")
print(f"  - 标签提取: {len(tags)} 个")
print(f"  - 信息匹配: {matched_count}/{len(tags)} 个成功")
print(f"  - 内容填充: {replacement_count} 处替换")
print(f"\n输出文件: {OUTPUT_PATH}")
