# -*- coding: utf-8 -*-
"""
完整功能测试 - 验证修复效果
"""

import os
import sys
import re
import json
import time
from pathlib import Path

if sys.platform == 'win32':
    import io
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8')

from docx import Document
from openai import OpenAI
from dotenv import load_dotenv

load_dotenv()

KIMI_API_KEY = os.getenv('KIMI_API_KEY', '')
KIMI_BASE_URL = os.getenv('KIMI_BASE_URL', 'https://api.moonshot.cn/v1')
KIMI_MODEL = 'kimi-k2.6'

TEMPLATE_PATH = r'E:\Users\han\Desktop\投标文件组成01.docx'
COMPANY_INFO_PATH = r'E:\Users\han\Desktop\企业投标信息01.txt'
OUTPUT_PATH = r'E:\Users\han\Desktop\auto_fill\outputs\test_output_v2.docx'

print("=" * 60)
print("完整功能测试 - 验证修复效果")
print(f"使用模型: {KIMI_MODEL}")
print("=" * 60)

client = OpenAI(api_key=KIMI_API_KEY, base_url=KIMI_BASE_URL)

# 加载文档
doc = Document(TEMPLATE_PATH)
print(f"文档: {len(doc.paragraphs)} 段落, {len(doc.tables)} 表格")

with open(COMPANY_INFO_PATH, 'r', encoding='utf-8') as f:
    company_info = f.read()
print(f"企业信息: {len(company_info)} 字符")

# Step 1: 段落打标
print()
print("【Step 1】段落打标")
print("-" * 60)

valid_paragraphs = [(i, p) for i, p in enumerate(doc.paragraphs) if p.text.strip()]
all_tags = set()
paragraph_processed = 0

for batch_start in range(0, len(valid_paragraphs), 10):
    batch_paras = valid_paragraphs[batch_start:batch_start + 10]
    text_block = "\n".join([f"[{idx}] {p.text.strip()}" for idx, p in batch_paras])

    prompt = f"""【角色】投标助理，识别空白处并打标签。

【空白处类型】
1. 下划线:________
2. 括号提示:(项目名称)、（单位盖章）
3. 散落空格:  年  月  日
4. 冒号留白:投标人:

【打标规则】
1. 标签格式: {{{{标签名}}}}
2. 标签必须具体:
   - 错误: {{{{姓名}}}}、{{{{地址}}}}、{{{{年}}}}
   - 正确: {{{{法定代表人姓名}}}}、{{{{投标人地址}}}}、{{{{投标日期年}}}}
3. 保留原文其他内容
4. 保留 [数字] 索引

【待打标文本】
{text_block}"""

    try:
        response = client.chat.completions.create(
            model=KIMI_MODEL,
            messages=[{"role": "user", "content": prompt}],
            max_tokens=2000
        )
        result = response.choices[0].message.content

        pattern = r"^\[(\d+)\]\s*(.*)$"
        for line in result.splitlines():
            m = re.match(pattern, line.strip())
            if m:
                idx = int(m.group(1))
                content = m.group(2)
                for orig_idx, para in batch_paras:
                    if orig_idx == idx and "{{" in content:
                        para.clear()
                        para.add_run(content)
                        paragraph_processed += 1
                        tags = re.findall(r"\{\{([^}]+)\}\}", content)
                        all_tags.update(tags)

        print(f"  批次 {batch_start//10 + 1}: OK")
        time.sleep(0.3)

    except Exception as e:
        print(f"  批次 {batch_start//10 + 1} 失败: {e}")

print(f"段落打标完成: {paragraph_processed} 个")
print(f"提取到标签: {len(all_tags)} 个")
print(f"标签示例: {sorted(list(all_tags))[:15]}")

# Step 2: 信息匹配
print()
print("【Step 2】信息匹配")
print("-" * 60)

tags_list = sorted(list(all_tags))
matched_data = {}

if tags_list:
    prompt = f"""【角色】企业信息匹配专家

【标签列表】
{json.dumps(tags_list, ensure_ascii=False, indent=2)}

【企业信息】
{company_info[:5000]}

【匹配规则】
1. 精确匹配标签语义
2. 日期类标签提取数字部分
3. 找不到返回"待补全"

【输出格式】
JSON对象，如:
{{"投标人单位名称": "江苏华宇市政建设集团有限公司"}}"""

    try:
        response = client.chat.completions.create(
            model=KIMI_MODEL,
            messages=[{"role": "user", "content": prompt}],
            max_tokens=4000
        )
        result = response.choices[0].message.content

        result = re.sub(r"^```json\s*", "", result)
        result = re.sub(r"^```\s*", "", result)
        result = re.sub(r"```\s*$", "", result)

        json_match = re.search(r"\{[\s\S]*\}", result)
        if json_match:
            matched_data = json.loads(json_match.group())
            for tag in tags_list:
                if tag not in matched_data:
                    matched_data[tag] = "待补全"
    except Exception as e:
        print(f"匹配失败: {e}")
        matched_data = {tag: "待补全" for tag in tags_list}

filled = sum(1 for v in matched_data.values() if v and v != "待补全")
total = len(matched_data)
print(f"匹配结果: {filled}/{total} 成功 ({filled*100//total if total else 0}%)")

print()
print("匹配结果预览:")
for i, (k, v) in enumerate(list(matched_data.items())[:20]):
    v_preview = str(v)[:40] if isinstance(v, str) else str(v)[:60]
    print(f"  {k}: {v_preview}")

# Step 3: 填充文档
print()
print("【Step 3】填充文档")
print("-" * 60)

fill_count = 0
for para in doc.paragraphs:
    for tag, value in matched_data.items():
        placeholder = f"{{{{{tag}}}}}"
        if placeholder in para.text:
            full_text = para.text
            new_text = full_text.replace(placeholder, str(value))
            if new_text != full_text:
                para.clear()
                para.add_run(new_text)
                fill_count += 1

print(f"填充完成: {fill_count} 处")

# Step 4: 保存
print()
print("【Step 4】保存结果")
print("-" * 60)

Path("outputs").mkdir(exist_ok=True)
doc.save(OUTPUT_PATH)
print(f"输出保存到: {OUTPUT_PATH}")

# 最终统计
print()
print("=" * 60)
print("测试结果总结")
print("=" * 60)
print(f"段落打标: {paragraph_processed} 个")
print(f"标签提取: {len(all_tags)} 个")
print(f"信息匹配: {filled}/{total} 成功 ({filled*100//total if total else 0}%)")
print(f"文档填充: {fill_count} 处")
